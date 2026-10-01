"""Bounded, local-only CSV/XLSX parsing. Never evaluate formulas or call AI."""

import csv
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

from defusedxml import ElementTree
from openpyxl import load_workbook
from pydantic import ValidationError

from app.errors import DomainError, fail
from app.roster_schemas import RosterIssue, RosterOptions, RosterRow

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 16
MAX_EXPANDED_BYTES = 20 * 1024 * 1024
HEADERS = {
    "studentnumber": "student_number",
    "studentno": "student_number",
    "studentid": "student_number",
    "studentemail": "email",
    "email": "email",
    "studentname": "display_name",
    "name": "display_name",
    "fullname": "display_name",
    "studentgradelevel": "grade_level",
    "gradelevel": "grade_level",
    "grade": "grade_level",
    "yearlevel": "grade_level",
    "section": "section",
}
REQUIRED = {"student_number", "email", "display_name", "grade_level"}


def clean_label(value):
    return " ".join(value.split())


def grade_label(value):
    value = clean_label(value)
    match = re.fullmatch(r"(?:grade\s*|g\s*)?(\d{1,2})", value, re.IGNORECASE)
    if match:
        number = int(match[1])
        if 1 <= number <= 12:
            return f"Grade {number}"
        raise ValueError("Use Grade 1-12 or an explicit year-level label")
    match = re.fullmatch(r"(?:college\s+)?year\s*(\d)|([1-6])(?:st|nd|rd|th)\s+year", value, re.IGNORECASE)
    if match:
        number = int(match[1] or match[2])
        if 1 <= number <= 6:
            return f"Year {number}"
        raise ValueError("Use Year 1-6")
    if value.casefold() in {"k", "kinder", "kindergarten"}:
        return "Kindergarten"
    # Other school-specific categories are retained, never inferred from age/email/number.
    return value


def fingerprint(data, filename, options):
    envelope = {"version": 1, "extension": Path(filename).suffix.lower(), **options.model_dump(mode="json")}
    return hashlib.sha256(json.dumps(envelope, sort_keys=True).encode() + b"\0" + data).hexdigest()


def _xlsx_rows(data, sheet_name):
    # Check expanded sizes and scan XML with entity expansion disabled before openpyxl.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 256 or sum(x.file_size for x in entries) > MAX_EXPANDED_BYTES:
            fail(413, "roster_too_large", "Workbook exceeds expanded size or archive entry limit")
        if len({x.filename for x in entries}) != len(entries):
            fail(422, "invalid_workbook", "Workbook has duplicate archive entries")
        nodes = 0
        for entry in entries:
            name = entry.filename.lower()
            if entry.flag_bits & 1 or "vbaproject" in name or name.startswith("xl/externallinks/"):
                fail(
                    415,
                    "unsupported_workbook",
                    "Encrypted, macro-enabled or externally linked workbooks are unsupported",
                )
            if name.endswith((".xml", ".rels")):
                depth = 0
                with archive.open(entry) as stream:
                    for event, node in ElementTree.iterparse(
                        stream, events=("start", "end"), forbid_dtd=True
                    ):
                        if event == "end":
                            node.clear()
                            depth -= 1
                            continue
                        nodes += 1
                        depth += 1
                        if nodes > 100_000 or depth > 64:
                            fail(413, "roster_too_large", "Workbook XML exceeds complexity limits")
                        if (
                            name.startswith("xl/worksheets/")
                            and node.tag == "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"
                        ):
                            match = re.fullmatch(r"([A-Z]+)([0-9]+)", node.attrib.get("r", ""))
                            if (
                                not match
                                or len(match[1]) != 1
                                or match[1] > "P"
                                or int(match[2]) > MAX_ROWS + 1
                            ):
                                fail(413, "roster_too_large", "Use at most 2,000 data rows and 16 columns")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if sheet_name is None and len(workbook.sheetnames) != 1:
            fail(422, "sheet_required", "Specify sheet_name for a workbook with multiple sheets")
        if sheet_name is not None and sheet_name not in workbook.sheetnames:
            fail(422, "sheet_not_found", "Requested worksheet was not found")
        sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        sheet.reset_dimensions()  # Do not trust stale Excel used-range metadata.
        for cells in sheet.iter_rows(max_col=MAX_COLUMNS, max_row=MAX_ROWS + 1):
            yield [(cell.value, cell.data_type == "f") for cell in cells]
    finally:
        workbook.close()


def parse_roster(data: bytes, filename: str, options: RosterOptions):
    if len(data) > MAX_BYTES:
        fail(413, "roster_too_large", "Roster file must be at most 5 MiB")
    extension = Path(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        fail(415, "unsupported_roster", "Upload UTF-8 CSV or Excel .xlsx; save legacy .xls as .xlsx")
    issues, rows, seen = [], [], set()

    def issue(number, field, code, message):
        issues.append(RosterIssue(row=number, field=field, code=code, message=message))

    try:
        if extension == ".csv":
            if options.sheet_name:
                fail(422, "unexpected_sheet", "sheet_name applies only to Excel workbooks")
            raw = csv.reader(io.StringIO(data.decode("utf-8-sig"), newline=""), strict=True)
            source = ([(value, False) for value in row] for row in raw)
        else:
            source = _xlsx_rows(data, options.sheet_name)
        headers = None
        total = 0
        for number, cells in enumerate(source, 1):
            if number > MAX_ROWS + 1 or len(cells) > MAX_COLUMNS:
                fail(413, "roster_too_large", "Use at most 2,000 data rows and 16 columns")
            # Excel reader pads to the configured column cap; trim unused trailing cells.
            while cells and cells[-1][0] in (None, ""):
                cells.pop()
            if number == 1:
                headers = []
                for value, formula in cells:
                    name = HEADERS.get(re.sub(r"[\s_\-]+", "", str(value)).casefold())
                    if formula or name is None or name in headers:
                        fail(
                            422,
                            "invalid_headers",
                            "Use unique supported column headers from the roster template",
                        )
                    headers.append(name)
                if not REQUIRED.issubset(headers):
                    fail(
                        422,
                        "missing_headers",
                        "Required columns: student_number, student_email, student_name, grade_level",
                    )
                continue
            if not cells:
                continue
            total += 1
            if len(cells) > len(headers):
                issue(number, "row", "extra_cells", "Row contains data outside the named columns")
                continue
            values = {"row": number, "section": options.default_section}
            before = len(issues)
            for field, (value, formula) in zip(headers, cells):
                if formula or (isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@"))):
                    issue(number, field, "formula_not_allowed", "Use plain values; formulas are not accepted")
                    continue
                if field == "student_number" and not isinstance(value, str):
                    issue(
                        number,
                        field,
                        "student_number_must_be_text",
                        "Format Excel student numbers as Text to preserve leading zeros",
                    )
                    continue
                if field == "grade_level" and isinstance(value, int) and not isinstance(value, bool):
                    value = str(value)
                if value is None or value == "":
                    if field != "section":
                        values[field] = ""
                    continue
                if not isinstance(value, str) or len(value) > 254 or any(ord(c) < 32 for c in value):
                    issue(
                        number,
                        field,
                        "invalid_cell",
                        "Use plain text within the field length limit, without control characters",
                    )
                    continue
                values[field] = value.strip()
            try:
                if "grade_level" in values:
                    values["grade_level"] = grade_label(values["grade_level"])
                for field in ("display_name", "section"):
                    if field in values:
                        values[field] = clean_label(values[field])
                row = RosterRow.model_validate(values)
            except ValidationError as exc:
                for error in exc.errors():
                    issue(
                        number, str(error["loc"][0]), "invalid_value", "Required value is missing or invalid"
                    )
                continue
            except ValueError as exc:
                issue(number, "grade_level", "invalid_grade_level", str(exc))
                continue
            if row.student_number in seen:
                issue(
                    number,
                    "student_number",
                    "duplicate_student_number",
                    "Student number occurs more than once in this file",
                )
            seen.add(row.student_number)
            if len(issues) == before:
                rows.append(row)
        if not headers or not total:
            fail(422, "empty_roster", "Roster requires a header and at least one student row")
        return rows, issues, total
    except DomainError:
        raise
    except Exception:
        # Provider/file exceptions can contain student values; never echo/log their details.
        fail(
            422,
            "invalid_roster_file",
            "Cannot read roster; use a valid UTF-8 CSV or unencrypted .xlsx workbook",
        )
